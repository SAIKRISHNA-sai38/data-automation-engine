from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.chart import BarChart, Reference


def export_to_excel(
    df,
    average_price,
    highest_price,
    lowest_price,
    total_records,
    valid_records,
    invalid_records,
    filename="product_report.xlsx"
):

    # Write product data
    df.to_excel(filename, index=False, sheet_name="Products")

    # Open workbook
    workbook = load_workbook(filename)
    sheet = workbook["Products"]

    # -----------------------------------
    # PRODUCTS SHEET
    # -----------------------------------

    for cell in sheet[1]:
        cell.font = Font(bold=True, size=12)
        cell.alignment = Alignment(horizontal="center")

    for row in range(2, len(df) + 2):
        sheet.cell(row, 2).number_format = '#,##0.00'

    # Borders
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin")
    )

    for row in sheet.iter_rows():
        for cell in row:
            if cell.value is not None:
                cell.border = thin_border
                cell.alignment = Alignment(vertical="center")

    sheet.column_dimensions["A"].width = 30
    sheet.column_dimensions["B"].width = 20

    # Freeze header
    sheet.freeze_panes = "A2"

    # -----------------------------------
    # DASHBOARD SHEET
    # -----------------------------------

    dashboard = workbook.create_sheet("Dashboard")

    dashboard["A1"] = "DATA AUTOMATION DASHBOARD"
    dashboard["A1"].font = Font(bold=True, size=18)

    dashboard["A3"] = "DATA SUMMARY"
    dashboard["A3"].font = Font(bold=True, size=14)

    summary = [
        ("Total Records", total_records),
        ("Valid Records", valid_records),
        ("Invalid Records", invalid_records),
        ("Average Price", average_price),
        ("Highest Price", highest_price),
        ("Lowest Price", lowest_price)
    ]

    for row, (label, value) in enumerate(summary, start=4):
        dashboard.cell(row, 1, label)
        dashboard.cell(row, 2, value)

    # Format prices
    for row in range(7, 10):
        dashboard.cell(row, 2).number_format = '#,##0.00'

    # Borders
    for row in range(4, 10):
        dashboard.cell(row, 1).border = thin_border
        dashboard.cell(row, 2).border = thin_border

    # -----------------------------------
    # CHART DATA
    # -----------------------------------

    chart_start = 12

    dashboard.cell(chart_start, 1, "Product")
    dashboard.cell(chart_start, 2, "Price")

    dashboard.cell(chart_start, 1).font = Font(bold=True)
    dashboard.cell(chart_start, 2).font = Font(bold=True)

    for index, row in df.iterrows():
        dashboard.cell(chart_start + index + 1, 1, row["name"])
        dashboard.cell(chart_start + index + 1, 2, row["price"])

    # -----------------------------------
    # CREATE BAR CHART
    # -----------------------------------

    chart = BarChart()

    chart.title = "Product Price Comparison"
    chart.y_axis.title = "Price"
    chart.x_axis.title = "Product"

    data = Reference(
        dashboard,
        min_col=2,
        min_row=chart_start,
        max_row=chart_start + len(df)
    )

    categories = Reference(
        dashboard,
        min_col=1,
        min_row=chart_start + 1,
        max_row=chart_start + len(df)
    )

    chart.add_data(data, titles_from_data=True)
    chart.set_categories(categories)

    chart.height = 8
    chart.width = 15

    dashboard.add_chart(chart, "D3")

    # -----------------------------------
    # DASHBOARD FORMATTING
    # -----------------------------------

    dashboard.column_dimensions["A"].width = 25
    dashboard.column_dimensions["B"].width = 20

    for row in dashboard.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="center")

    # Save
    workbook.save(filename)

    print(f"Excel report created: {filename}")